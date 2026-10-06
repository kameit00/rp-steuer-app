"""Synthetic regression tests against the untouched, direct production functions.

Optional frozen integration:
    RP_STEUER_APP_TEST_ENGINE=/absolute/path/to/rp-steuer-app-engine python -m unittest discover -s tests -v
"""
from __future__ import annotations

import contextlib
import csv
import hashlib
import io
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest import mock

import openpyxl

import engine
from compare_workbooks import assert_workbooks_equal
from rp_steuer.sevdesk_csv_to_excel import convert_csv_to_excel
from rp_steuer.sevdesk_paypal import process_csv_files

ROOT = Path(__file__).resolve().parents[1]
PRODUCTION = ROOT / "production"
FROZEN_ENGINE = os.environ.get("RP_STEUER_APP_TEST_ENGINE")
PAYPAL_HEADERS = ["Bezahldatum", "Beschreibung", "Name", "Betrag", "", ""]
PAYPAL_ROWS = [
    ["01.01.2026", "Kundenkauf und Gebühren", "Müller", "€ 4.000,00", "", ""],
    ["2026-01-03", "Gebühren", "PayPal", "-12,34 €", "", ""],
    ["31.01.2026", "Ohne Zuordnung", "Zoë", "100.50", "", ""],
    ["05.02.2026", "Rückerstattung", "Müller", "-20,00", "", ""],
    ["2026-02-06", "Ausgleich", "Bankeinzug", "1.234,56€", "", ""],
    ["kein Datum", "Gebühren", "PayPal", "-1,25", "", ""],
    ["", "Ohne Zuordnung", "François", "3,50", "", ""],
    ["31.02.2026", "Ohne Zuordnung", "Åsa", "nicht numerisch", "", ""],
]
CUSTOM_RULES = {
    "rules": [
        {"column": "Beschreibung", "substring": "Kundenkauf", "info_column_value": "Kunde"},
        {"column": "Beschreibung", "substring": "Rückerstattung", "info_column_value": "Kunde"},
        {"column": "Beschreibung", "substring": "Gebühren", "info_column_value": "Gebühr"},
        {"column": "Name", "substring": "Bankeinzug", "info_column_value": "Bank"},
    ],
    "always_show_in_summary": ["Immer leer", "Kunde"],
}
INVOICE_HEADERS = ["Rechnungsnummer", "Rechnungs-Datum", "Fällig am", "Bezahlt am", "Kunde", "Gesamtbetrag-Netto", "Gesamtbetrag-Brutto", "offener Betrag", "", ""]
INVOICE_ROWS = [
    ["RE-001", "05.01.2026", "19.01.2026", "12.01.2026", "Müller", "100,25", "119,30", "0,00", "", ""],
    ["RE-002", "17.01.2026", "31.01.2026", "01.02.2026", "Zoë", "200,50", "238,60", "0,00", "", ""],
    ["RE-003", "30.01.2026", "13.02.2026", "", "Åsa", "50,00", "59,50", "59,50", "", ""],
    ["RE-004", "07.02.2026", "21.02.2026", "09.02.2026", "François", "10,00", "11,90", "0,00", "", ""],
    ["RE-005", "ungültig", "ungültig", "ungültig", "No date", "3,75", "4,46", "4,46", "", ""],
]
ON_TIME_ROWS = [INVOICE_ROWS[0], INVOICE_ROWS[3]]


def write_csv(path, headers, rows, delimiter=";"):
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream, delimiter=delimiter)
        writer.writerow(headers)
        writer.writerows(rows)
    return path


class EngineTestCase(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory(prefix="rp-tests-ä-税-")
        self.addCleanup(temporary.cleanup)
        # macOS /var is a symlink to /private/var; the engine resolves paths.
        self.directory = Path(temporary.name).resolve()
        self.inputs = self.directory / "Eingänge ü 测试"
        self.output = self.directory / "Ausgabe ö 测试"
        self.inputs.mkdir()
        self.output.mkdir()
        self.rules = self.inputs / "Regeln ü.json"
        self.rules.write_text(json.dumps(CUSTOM_RULES, ensure_ascii=False), encoding="utf-8")
        self.events = []

    def paypal(self, name="PayPal ü.csv", rows=None, delimiter=";"):
        return write_csv(self.inputs / name, PAYPAL_HEADERS, PAYPAL_ROWS if rows is None else rows, delimiter)

    def invoice(self, name="Rechnungen ü.csv", rows=None):
        return write_csv(self.inputs / name, INVOICE_HEADERS, INVOICE_ROWS if rows is None else rows)

    def request(self, mode, sources, rules=True):
        request = {"mode": mode, "inputs": [str(source) for source in sources], "output_directory": str(self.output)}
        if mode == "paypal" and rules:
            request["rules_file"] = str(self.rules)
        return request

    def execute(self, request):
        return engine.run_job(request, report=self.events.append)

    def baseline(self, mode, source, rules=None):
        work = self.directory / f"baseline-{len(list(self.directory.glob('baseline-*')))}"
        work.mkdir()
        copied = work / (source.stem + ".csv")
        shutil.copyfile(source, copied)
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            if mode == "paypal":
                shutil.copyfile(rules or PRODUCTION / "rules.json", work / "rules.json")
                process_csv_files(delete_csv=False, directory=str(work))
            else:
                self.assertIsNotNone(convert_csv_to_excel(str(copied)))
        return {path.name: path for path in work.glob("*.xlsx")}

    def compare_with_baseline(self, mode, source, rules=True):
        before = source.read_bytes()
        expected = self.baseline(mode, source, self.rules if rules else None)
        self.assertTrue(expected, "Fixture must generate real production workbooks")
        files = self.execute(self.request(mode, [source], rules=rules))
        actual = {Path(path).name: Path(path) for path in files}
        self.assertEqual(set(actual), set(expected))
        self.assertEqual({path.name for path in self.output.iterdir()}, set(expected))
        for name, path in actual.items():
            self.assertEqual(path.parent, self.output)
            assert_workbooks_equal(expected[name], path)
        self.assertEqual(source.read_bytes(), before, "Input CSV must remain untouched")
        return actual

    def load_workbook(self, path):
        workbook = openpyxl.load_workbook(path, data_only=False)
        self.addCleanup(workbook.close)
        return workbook

    def assert_failure_unchanged(self, request, error_type=Exception):
        before_inputs = {path: path.read_bytes() for path in self.inputs.rglob("*") if path.is_file()}
        before_outputs = {path: path.read_bytes() for path in self.output.iterdir() if path.is_file()}
        with self.assertRaises(error_type):
            self.execute(request)
        self.assertEqual({path: path.read_bytes() for path in self.inputs.rglob("*") if path.is_file()}, before_inputs)
        self.assertEqual({path: path.read_bytes() for path in self.output.iterdir() if path.is_file()}, before_outputs)

    def protocol(self, payload, command=None, raw=False):
        request_path = self.directory / "Auftrag ü 税.json"
        request_path.write_text(payload if raw else json.dumps(payload, ensure_ascii=False), encoding="utf-8")
        result = subprocess.run((command or [sys.executable, str(ROOT / "engine.py")]) + ["--request", str(request_path)], cwd=self.directory, capture_output=True, text=True, encoding="utf-8", timeout=120)
        events = [json.loads(line) for line in result.stdout.splitlines()]
        self.assertTrue(events, result.stderr)
        self.assertTrue(all(isinstance(event, dict) and "type" in event for event in events))
        final = [event for event in events if event["type"] == "result"]
        self.assertEqual(len(final), 1, events)
        self.assertEqual(events[-1], final[0], "Result must be the final protocol event")
        self.assertEqual(result.stderr, "", result.stderr)
        return result, events, final[0]


class EquivalenceCases:
    def test_paypal_euro_thousands_unknown_dates_categories_and_multiple_months(self):
        actual = self.compare_with_baseline("paypal", self.paypal())
        self.assertEqual(set(actual), {"2026-01-Paypal.xlsx", "2026-02-Paypal.xlsx", "unknown-date-Paypal.xlsx"})
        january = self.load_workbook(actual["2026-01-Paypal.xlsx"]).active
        self.assertEqual(january.title, "PayPal Transactions")
        self.assertEqual([cell.value for cell in january[1]], ["Bezahldatum", "Beschreibung", "Name", "Betrag", "Info"])
        self.assertEqual(january["D2"].value, 4000)
        self.assertEqual(january["D3"].value, -12.34)
        self.assertEqual(january["D4"].value, 100.5)
        self.assertEqual(january["E2"].value, "Kunde", "First matching rule wins")
        self.assertEqual(january["D2"].number_format, "#,##0.00 €")
        labels = {row[0].value: row[1].value for row in january.iter_rows()}
        self.assertEqual(labels["'Betrag' gesamt:"], "=SUM(D2:D4)")
        self.assertIn('"Immer leer"', labels["Info 'Immer leer':"])
        self.assertIn('""', labels["Info nicht ausgefüllt:"])
        self.assertTrue(labels["Gegenprüfung muss immer 0 sein:"].startswith("=B"))
        self.assertTrue(january["A1"].font.bold)
        self.assertGreater(january.column_dimensions["A"].width, 0)
        february = self.load_workbook(actual["2026-02-Paypal.xlsx"]).active
        self.assertEqual(february["D3"].value, 1234.56)
        unknown = self.load_workbook(actual["unknown-date-Paypal.xlsx"]).active
        self.assertEqual(unknown["D4"].value, "nicht numerisch")

    def test_paypal_legacy_list_rules_and_comma_delimiter(self):
        self.rules.write_text(json.dumps(CUSTOM_RULES["rules"], ensure_ascii=False), encoding="utf-8")
        actual = self.compare_with_baseline("paypal", self.paypal(delimiter=","))
        january = self.load_workbook(actual["2026-01-Paypal.xlsx"]).active
        self.assertNotIn("Info 'Immer leer':", [row[0].value for row in january.iter_rows()])

    def test_paypal_default_production_rules(self):
        self.compare_with_baseline("paypal", self.paypal(), rules=False)

    def test_invoice_paid_late_unpaid_mixed_month_dates_decimal_and_empty_columns(self):
        actual = self.compare_with_baseline("sevdesk", self.invoice())
        self.assertEqual(set(actual), {"2026-01-Rechnungen.xlsx"})
        workbook = self.load_workbook(actual["2026-01-Rechnungen.xlsx"])
        self.assertEqual(workbook.sheetnames, ["Alle Rechnungen", "Verspätete Zahlungen"])
        all_invoices = workbook.worksheets[0]
        self.assertEqual(all_invoices.max_column, 8, "Trailing empty CSV columns must be removed")
        self.assertEqual(all_invoices["F2"].value, 100.25)
        self.assertEqual(all_invoices["G2"].value, 119.3)
        self.assertEqual(all_invoices["F7"].value, "=SUM(F2:F6)")
        self.assertEqual(all_invoices["B2"].value.strftime("%Y-%m-%d"), "2026-01-05")
        self.assertEqual(all_invoices["B2"].number_format, "dd.mm.yyyy")
        self.assertEqual(all_invoices["F2"].number_format, "#,##0.00 €")
        self.assertTrue(all_invoices["F7"].font.bold)
        self.assertEqual(all_invoices["B6"].value, None, "Invalid dates are coerced to empty cells")
        late = workbook.worksheets[1]
        self.assertEqual([late.cell(row, 1).value for row in range(2, 5)], ["RE-002", "RE-003", "RE-005"])
        self.assertEqual(late["F5"].value, "=SUM(F2:F4)")

    def test_invoice_no_late_payments_has_empty_second_sheet(self):
        actual = self.compare_with_baseline("sevdesk", self.invoice(rows=ON_TIME_ROWS))
        workbook = self.load_workbook(next(iter(actual.values())))
        self.assertEqual(workbook.sheetnames, ["Alle Rechnungen", "Verspätete Zahlungen"])
        self.assertEqual(list(workbook.worksheets[1].values), [])

    def test_uppercase_csv_extensions_in_both_modes(self):
        for mode in ("paypal", "sevdesk"):
            with self.subTest(mode=mode):
                for path in self.output.iterdir():
                    path.unlink()
                source = self.paypal("PayPal.CSV") if mode == "paypal" else self.invoice("Invoices.CSV")
                self.compare_with_baseline(mode, source)

    def test_invoice_all_invalid_dates_uses_input_stem(self):
        row = list(INVOICE_ROWS[-1])
        actual = self.compare_with_baseline("sevdesk", self.invoice("Ohne Datum ü.csv", rows=[row]))
        self.assertEqual(set(actual), {"Ohne Datum ü.xlsx"})


class WrapperEquivalenceTests(EquivalenceCases, EngineTestCase):
    pass


@unittest.skipUnless(FROZEN_ENGINE, "Set RP_STEUER_APP_TEST_ENGINE to test the frozen executable")
class FrozenEngineEquivalenceTests(EquivalenceCases, EngineTestCase):
    def execute(self, request):
        executable = str(Path(FROZEN_ENGINE).expanduser().resolve())
        result, events, final = self.protocol(request, command=[executable])
        self.events.extend(events)
        self.assertEqual(result.returncode, 0, final)
        self.assertIs(final["success"], True)
        return final["files"]


class WrapperSafetyTests(EngineTestCase):
    def test_success_never_deletes_inputs_rules_or_unrelated_output_files(self):
        source = self.paypal()
        original = source.read_bytes()
        rules = self.rules.read_bytes()
        sentinel = self.output / "unrelated.txt"
        sentinel.write_bytes(b"do not delete")
        files = self.execute(self.request("paypal", [source]))
        self.assertEqual(len(files), 3)
        self.assertEqual(source.read_bytes(), original)
        self.assertEqual(self.rules.read_bytes(), rules)
        self.assertEqual(sentinel.read_bytes(), b"do not delete")

    def test_never_overwrites_existing_output_case_insensitively_for_both_modes(self):
        for mode, name in [("paypal", "2026-02-paypal.XLSX"), ("sevdesk", "2026-01-rechnungen.XLSX")]:
            with self.subTest(mode=mode):
                sentinel = self.output / name
                sentinel.write_bytes(b"existing workbook must not be touched")
                source = self.paypal() if mode == "paypal" else self.invoice()
                self.assert_failure_unchanged(self.request(mode, [source]), engine.JobError)
                sentinel.unlink()

    def test_duplicate_output_names_fail_whole_job_for_both_modes(self):
        for mode in ("paypal", "sevdesk"):
            with self.subTest(mode=mode):
                sources = [self.paypal("one.csv"), self.paypal("two.csv")] if mode == "paypal" else [self.invoice("one.csv"), self.invoice("two.csv")]
                self.assert_failure_unchanged(self.request(mode, sources), engine.JobError)
                self.assertEqual(list(self.output.iterdir()), [])

    def test_unknown_date_output_collision_also_fails_whole_job(self):
        first = self.paypal("first.csv", rows=[PAYPAL_ROWS[-1]])
        second = self.paypal("second.csv", rows=[PAYPAL_ROWS[-2]])
        self.assert_failure_unchanged(self.request("paypal", [first, second]), engine.JobError)

    def test_identical_input_selected_twice_is_rejected(self):
        source = self.paypal()
        self.assert_failure_unchanged(self.request("paypal", [source, source]), engine.JobError)

    def test_later_input_failure_discards_already_staged_paypal_and_invoice_outputs(self):
        for mode in ("paypal", "sevdesk"):
            with self.subTest(mode=mode):
                valid = self.paypal("valid.csv") if mode == "paypal" else self.invoice("valid.csv")
                invalid = self.inputs / "empty.csv"
                invalid.write_bytes(b"")
                self.assert_failure_unchanged(self.request(mode, [valid, invalid]), engine.JobError)
                self.assertEqual(list(self.output.iterdir()), [])

    def test_invoice_processing_error_from_missing_required_columns(self):
        source = write_csv(self.inputs / "broken.csv", ["Rechnungs-Datum", "Gesamtbetrag-Netto"], [["01.03.2026", "1,00"]])
        self.assert_failure_unchanged(self.request("sevdesk", [source]), engine.JobError)
        self.assertTrue(any("Error creating Excel file" in event.get("message", "") for event in self.events))

    def test_invoice_bad_csv_encoding_and_empty_csv_raise_without_final_writes(self):
        for payload in (b"", b"\xff\xfeinvalid"):
            with self.subTest(payload=payload):
                source = self.inputs / "broken.csv"
                source.write_bytes(payload)
                self.assert_failure_unchanged(self.request("sevdesk", [source]), engine.JobError)

    def test_paypal_malformed_empty_header_only_or_ignored_csv_raise_without_final_writes(self):
        cases = [("broken.csv", b"not a CSV"), ("empty.csv", b""), ("header.csv", b"Bezahldatum;Betrag;Beschreibung\n"), ("2026-01-Paypal.csv", b"Bezahldatum;Betrag\n01.01.2026;1,00\n"), ("encoding.csv", b"\xff\xfeinvalid")]
        for name, payload in cases:
            with self.subTest(name=name):
                source = self.inputs / name
                source.write_bytes(payload)
                self.assert_failure_unchanged(self.request("paypal", [source]), engine.JobError)

    def test_invalid_rule_shapes_are_rejected_before_processing(self):
        source = self.paypal()
        invalid_rules = [None, {}, "rules", 1, {"rules": {}}, {"rules": [None]}, {"rules": [{"column": "Beschreibung", "substring": "x"}]}, {"rules": [{"column": "Beschreibung", "substring": 1, "info_column_value": "x"}]}, {"rules": [], "always_show_in_summary": "x"}, {"rules": [], "always_show_in_summary": [1]}]
        for rules in invalid_rules:
            with self.subTest(rules=rules):
                self.rules.write_text(json.dumps(rules), encoding="utf-8")
                with mock.patch.object(engine, "process_csv_files") as process:
                    self.assert_failure_unchanged(self.request("paypal", [source]), engine.JobError)
                    process.assert_not_called()

    def test_missing_or_malformed_rule_file_raises_without_writes(self):
        source = self.paypal()
        self.rules.write_text("{not-json", encoding="utf-8")
        self.assert_failure_unchanged(self.request("paypal", [source]), json.JSONDecodeError)
        self.rules.unlink()
        self.assert_failure_unchanged(self.request("paypal", [source]), FileNotFoundError)

    def test_invalid_mode_and_malformed_request_fields(self):
        source = self.paypal()
        good = self.request("paypal", [source])
        bad_requests = [None, [], "text", 1, {}, *[dict(good, mode=value) for value in (None, "PAYPAL", "other", 1)], *[dict(good, inputs=value) for value in (None, [], "file.csv", [1], [str(source), None])], *[dict(good, output_directory=value) for value in (None, "", 1)], dict(good, rules_file=1)]
        for request in bad_requests:
            with self.subTest(request=request):
                self.assert_failure_unchanged(request, engine.JobError)

    def test_missing_non_csv_directory_input_and_missing_output(self):
        good = self.paypal()
        text = self.inputs / "input.txt"
        text.write_text("data", encoding="utf-8")
        for source in (self.inputs / "missing.csv", text, self.inputs):
            with self.subTest(source=source):
                self.assert_failure_unchanged(self.request("paypal", [source]), engine.JobError)
        for output in (self.directory / "does-not-exist", good):
            with self.subTest(output=output):
                request = self.request("paypal", [good])
                request["output_directory"] = str(output)
                self.assert_failure_unchanged(request, engine.JobError)

    def test_production_returns_without_generating_files_are_rejected(self):
        for mode, function in [("paypal", "process_csv_files"), ("sevdesk", "convert_csv_to_excel")]:
            with self.subTest(mode=mode):
                source = self.paypal() if mode == "paypal" else self.invoice()
                with mock.patch.object(engine, function, return_value=None):
                    self.assert_failure_unchanged(self.request(mode, [source]), engine.JobError)

    def test_printed_production_failure_rejects_a_generated_workbook(self):
        source = self.paypal()

        def failing_production(*args, **kwargs):
            process_csv_files(*args, **kwargs)
            print("Failed synthetic failure after generating workbooks")

        with mock.patch.object(engine, "process_csv_files", side_effect=failing_production):
            self.assert_failure_unchanged(self.request("paypal", [source]), engine.JobError)
        self.assertTrue(any("Failed synthetic" in event.get("message", "") for event in self.events))

    def test_publication_failure_rolls_back_previous_outputs(self):
        source = self.paypal()
        original_open = Path.open

        def failing_open(path, mode="r", *args, **kwargs):
            if mode == "xb" and path.parent == self.output and path.name == "2026-02-Paypal.xlsx":
                raise OSError("simulated destination failure")
            return original_open(path, mode, *args, **kwargs)

        with mock.patch.object(Path, "open", failing_open):
            self.assert_failure_unchanged(self.request("paypal", [source]), OSError)
        self.assertEqual(list(self.output.iterdir()), [])

    def test_exclusive_creation_does_not_overwrite_a_concurrent_file(self):
        source = self.paypal()
        destination = self.output / "2026-01-Paypal.xlsx"
        original_open = Path.open

        def raced_open(path, mode="r", *args, **kwargs):
            if mode == "xb" and path == destination:
                destination.write_bytes(b"other job won the race")
            return original_open(path, mode, *args, **kwargs)

        with mock.patch.object(Path, "open", raced_open):
            with self.assertRaises(FileExistsError):
                self.execute(self.request("paypal", [source]))
        self.assertEqual(destination.read_bytes(), b"other job won the race")
        self.assertEqual(list(self.output.iterdir()), [destination])
        self.assertTrue(source.is_file())


class SubprocessProtocolTests(EngineTestCase):
    def test_success_protocol_with_unicode_paths_matches_direct_production(self):
        for mode in ("paypal", "sevdesk"):
            with self.subTest(mode=mode):
                for path in self.output.iterdir():
                    path.unlink()
                source = self.paypal() if mode == "paypal" else self.invoice()
                original = source.read_bytes()
                expected = self.baseline(mode, source, self.rules)
                result, events, final = self.protocol(self.request(mode, [source]))
                self.assertEqual(result.returncode, 0, final)
                self.assertIs(final["success"], True)
                self.assertTrue(any(event["type"] == "log" for event in events))
                self.assertEqual({Path(path).name for path in final["files"]}, set(expected))
                for path in final["files"]:
                    assert_workbooks_equal(expected[Path(path).name], path)
                self.assertEqual(source.read_bytes(), original)

    def test_failure_protocol_nonzero_exit_no_outputs_and_error(self):
        for mode in ("paypal", "sevdesk"):
            with self.subTest(mode=mode):
                valid = self.paypal() if mode == "paypal" else self.invoice()
                invalid = self.inputs / "empty.csv"
                invalid.write_bytes(b"")
                result, events, final = self.protocol(self.request(mode, [valid, invalid]))
                self.assertEqual(result.returncode, 1, events)
                self.assertIs(final["success"], False)
                self.assertEqual(final["files"], [])
                self.assertTrue(final["error"])
                self.assertEqual(list(self.output.iterdir()), [])
                self.assertTrue(valid.is_file())

    def test_malformed_json_and_invalid_request_are_machine_readable_failures(self):
        for payload in ("{not-json", "null", "[]", '{"mode":"unknown"}'):
            with self.subTest(payload=payload):
                result, events, final = self.protocol(payload, raw=True)
                self.assertEqual(result.returncode, 1, events)
                self.assertIs(final["success"], False)
                self.assertEqual(final["files"], [])
                self.assertTrue(final["error"])
                self.assertEqual(list(self.output.iterdir()), [])

    def test_missing_request_file_emits_failure_result(self):
        result = subprocess.run([sys.executable, str(ROOT / "engine.py"), "--request", str(self.directory / "missing.json")], capture_output=True, text=True, timeout=120)
        self.assertEqual(result.returncode, 1)
        final = json.loads(result.stdout)
        self.assertEqual(final["type"], "result")
        self.assertIs(final["success"], False)
        self.assertEqual(final["files"], [])


class ProductionSnapshotTests(unittest.TestCase):
    def test_snapshot_hashes_cover_all_production_sources_rules_and_lock(self):
        manifest = json.loads((PRODUCTION / "SHA256SUMS.json").read_text(encoding="utf-8"))
        self.assertIsInstance(manifest, dict)
        required = {str(path.relative_to(PRODUCTION)) for path in (PRODUCTION / "rp_steuer").rglob("*.py")} | {"rules.json", "production-uv.lock"}
        self.assertTrue(required <= manifest.keys(), f"Missing snapshot hashes: {required - manifest.keys()}")
        for relative, expected in manifest.items():
            with self.subTest(path=relative):
                path = PRODUCTION / relative
                self.assertTrue(path.resolve().is_relative_to(PRODUCTION.resolve()))
                self.assertRegex(expected, r"^[0-9a-f]{64}$")
                self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), expected)

    def test_optional_downloads_sources_match_byte_for_byte(self):
        external = Path.home() / "Downloads" / "rp-steuer"
        if not (external / "src" / "rp_steuer").is_dir():
            self.skipTest("Original Downloads/rp-steuer source is not present")
        for source in (PRODUCTION / "rp_steuer").rglob("*.py"):
            with self.subTest(source=source.name):
                original = external / "src" / source.relative_to(PRODUCTION)
                self.assertTrue(original.is_file(), original)
                self.assertEqual(source.read_bytes(), original.read_bytes())
        for relative, original in [("rules.json", external / "rules.json"), ("production-uv.lock", external / "uv.lock")]:
            if original.is_file():
                with self.subTest(source=relative):
                    self.assertEqual((PRODUCTION / relative).read_bytes(), original.read_bytes())


if __name__ == "__main__":
    unittest.main()
